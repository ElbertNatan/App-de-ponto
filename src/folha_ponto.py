"""
Folha de Ponto — app de controle de horas trabalhadas.

Recursos:
  - UI moderna com abas: Dia, Mes atual, Banco geral, Marcacoes
  - Aba Dia com carrossel (setas) pra navegar entre dias e editar horarios
  - Edicao manual de Inicio / Ini Pausa / Fim Pausa / Fim de qualquer dia
  - Cronometro live: se editar horarios do dia atual, o contador recalcula na hora
  - Botoes Play / Pause / Finalizar registram horarios automaticamente no dia atual
  - Excel folha_ponto.xlsx: uma aba por mes + Resumo + Marcacoes
  - Mantem tambem os .txt mensais (Excel + .txt)
  - Migracao automatica dos ponto_YYYY_MM.txt na primeira execucao
  - Banco de horas: meta 8h em dias uteis; FERIADO/ATESTADO/FOLGA neutralizam
"""

import calendar
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime, timedelta, date, time as dtime
from pathlib import Path
import re

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

try:
    from PIL import Image, ImageDraw, ImageFilter, ImageTk
    PIL_OK = True
except ImportError:
    PIL_OK = False


def _hex_para_rgb(hex_color):
    h = hex_color.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ---------- Constantes ----------

SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR.parent / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
EXCEL_PATH = DATA_DIR / "folha_ponto.xlsx"
ICONE_PATH = SCRIPT_DIR / "folha_ponto.ico"
APP_USER_MODEL_ID = "folha.ponto.desktop.app"
META_DIARIA = timedelta(hours=8)
TIPOS_NEUTRALIZA_META = ("FERIADO", "ATESTADO", "FOLGA")
AUTOSAVE_INTERVALO_MS = 10 * 60 * 1000  # 10 min


def dias_uteis_no_mes(ano, mes):
    _, ultimo = calendar.monthrange(ano, mes)
    return sum(1 for d in range(1, ultimo + 1) if date(ano, mes, d).weekday() < 5)

SHEET_RESUMO = "Resumo"
SHEET_MARCACOES = "Marcacoes"

# Indices das colunas (1-based)
COL_DATA = 1
COL_DIA = 2
COL_INICIO = 3
COL_INI_PAUSA = 4
COL_FIM_PAUSA = 5
COL_FIM = 6
COL_TRABALHO = 7
COL_PAUSA = 8
COL_META = 9
COL_SALDO = 10
COL_TIPO = 11
COL_OBS = 12

HEADERS_DIA = [
    "Data", "Dia", "Inicio", "Ini Pausa", "Fim Pausa", "Fim",
    "Trabalho", "Pausa", "Meta", "Saldo", "Tipo", "Observacao",
]
HEADERS_MARC = ["Data", "Tipo", "Observacao"]
HEADERS_RESUMO = [
    "Mes", "Dias trabalhados", "Trabalho total",
    "Meta total", "Saldo do mes", "Saldo acumulado",
]

DIAS_SEMANA = ["Segunda", "Terca", "Quarta", "Quinta", "Sexta", "Sabado", "Domingo"]
TIPOS_DIA = ["NORMAL", "FERIADO", "ATESTADO", "FOLGA"]

MESES_PT = [
    "Janeiro", "Fevereiro", "Marco", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

# Paleta — slate + indigo, contraste suave
COR_BG = "#f1f5f9"           # slate-100
COR_CARD = "#ffffff"
COR_HOVER = "#f8fafc"        # slate-50
COR_PRIMARY = "#4f46e5"      # indigo-600
COR_PRIMARY_HOVER = "#4338ca"
COR_PRIMARY_LIGHT = "#e0e7ff"
COR_PRIMARY_TEXT = "#3730a3"
COR_SUCCESS = "#16a34a"      # green-600
COR_SUCCESS_HOVER = "#15803d"
COR_SUCCESS_LIGHT = "#dcfce7"
COR_DANGER = "#dc2626"       # red-600
COR_DANGER_HOVER = "#b91c1c"
COR_DANGER_LIGHT = "#fee2e2"
COR_WARNING = "#ea580c"      # orange-600
COR_WARNING_HOVER = "#c2410c"
COR_WARNING_LIGHT = "#ffedd5"
COR_NEUTRAL = "#475569"      # slate-600
COR_NEUTRAL_HOVER = "#334155"
COR_TEXT = "#0f172a"         # slate-900
COR_TEXT_SOFT = "#334155"    # slate-700
COR_MUTED = "#64748b"        # slate-500
COR_BORDER = "#e2e8f0"       # slate-200
COR_BORDER_FORTE = "#cbd5e1" # slate-300
COR_DISABLED = "#94a3b8"     # slate-400
COR_ZEBRA = "#f8fafc"

# Estilo Excel
if OPENPYXL_OK:
    EX_HEADER_FILL = PatternFill("solid", fgColor="4C6EF5")
    EX_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    EX_ALT_FILL = PatternFill("solid", fgColor="F1F3F5")
    EX_TOTAL_FILL = PatternFill("solid", fgColor="E7F5FF")
    EX_BORDER = Border(
        left=Side(style="thin", color="DEE2E6"),
        right=Side(style="thin", color="DEE2E6"),
        top=Side(style="thin", color="DEE2E6"),
        bottom=Side(style="thin", color="DEE2E6"),
    )
    EX_TOTAL_BORDER = Border(
        top=Side(style="medium", color="4C6EF5"),
        bottom=Side(style="thin", color="DEE2E6"),
        left=Side(style="thin", color="DEE2E6"),
        right=Side(style="thin", color="DEE2E6"),
    )


# ---------- Helpers ----------

def td_to_str(td):
    total = int(td.total_seconds())
    sinal = "-" if total < 0 else ""
    total = abs(total)
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{sinal}{h:02}:{m:02}:{s:02}"


def str_to_td(s):
    if s is None:
        return timedelta()
    s = str(s).strip()
    if not s or s == "-":
        return timedelta()
    sinal = -1 if s.startswith("-") else 1
    s = s.lstrip("-")
    parts = s.split(":")
    try:
        h = int(parts[0])
        m = int(parts[1]) if len(parts) > 1 else 0
        sec = int(parts[2]) if len(parts) > 2 else 0
        return sinal * timedelta(hours=h, minutes=m, seconds=sec)
    except (ValueError, IndexError):
        return timedelta()


def parse_data(s):
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    if not s:
        return None
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_hhmmss(s):
    """Retorna time() ou None se invalido/vazio."""
    if s is None:
        return None
    s = str(s).strip()
    if not s or s == "-":
        return None
    # Aceita H:MM, HH:MM ou HH:MM:SS
    parts = s.split(":")
    try:
        h = int(parts[0])
        m = int(parts[1]) if len(parts) > 1 else 0
        sec = int(parts[2]) if len(parts) > 2 else 0
        if not (0 <= h < 24 and 0 <= m < 60 and 0 <= sec < 60):
            return None
        return dtime(h, m, sec)
    except (ValueError, IndexError):
        return None


def time_to_str(t):
    if t is None:
        return "-"
    if isinstance(t, datetime):
        t = t.time()
    return t.strftime("%H:%M:%S")


def combinar(data, t):
    if not data or not t:
        return None
    if isinstance(t, datetime):
        return t
    return datetime.combine(data, t)


def calcular_trab_pausa(inicio, ini_pausa, fim_pausa, fim, agora=None):
    """A partir dos 4 datetime (ou None), retorna (trabalho, pausa).
    Se fim is None e `agora` informado, usa `agora` como fim virtual (live mode).
    """
    if not inicio:
        return timedelta(), timedelta()
    fim_efetivo = fim or agora
    if not fim_efetivo and not ini_pausa:
        return timedelta(), timedelta()
    if ini_pausa and not fim_pausa:
        # ainda em pausa
        trab = ini_pausa - inicio
        pausa_fim_virtual = agora if (agora and not fim) else fim
        if pausa_fim_virtual:
            pausa = pausa_fim_virtual - ini_pausa
        else:
            pausa = timedelta()
        return _clamp(trab), _clamp(pausa)
    if ini_pausa and fim_pausa and fim_efetivo:
        trab = (ini_pausa - inicio) + (fim_efetivo - fim_pausa)
        pausa = fim_pausa - ini_pausa
        return _clamp(trab), _clamp(pausa)
    if fim_efetivo:
        return _clamp(fim_efetivo - inicio), timedelta()
    return timedelta(), timedelta()


def _clamp(td):
    return td if td.total_seconds() > 0 else timedelta()


# ---------- Parsing dos .txt legados ----------

_RE_DATA = re.compile(r"Data:\s*(\S+)")
_RE_EVENTO = re.compile(r"\[(\d{2}:\d{2}:\d{2})\]\s*(.+)")
_RE_TRAB = re.compile(r"Trabalho total:\s*(-?[\d:]+)")
_RE_PAUSA = re.compile(r"Pausa total:\s*(-?[\d:]+)")


def _norm(s):
    repl = {"Í": "I", "í": "i", "Ç": "C", "ç": "c", "Ã": "A", "ã": "a"}
    out = s.upper()
    for k, v in repl.items():
        out = out.replace(k, v.upper())
    return out


def parse_blocos_txt(path):
    """Lista de dicts: {data, inicio, ini_pausa, fim_pausa, fim, trabalho, pausa}."""
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocos = re.split(r"={5,}\s*\n", text)
    out = []
    for bloco in blocos:
        bloco = bloco.strip()
        if not bloco:
            continue
        m = _RE_DATA.search(bloco)
        if not m:
            continue
        d = parse_data(m.group(1))
        if not d:
            continue
        inicio = ini_pausa = fim_pausa = fim = None
        for hora, desc in _RE_EVENTO.findall(bloco):
            desc_n = _norm(desc)
            t = datetime.combine(d, datetime.strptime(hora, "%H:%M:%S").time())
            if "INICIO DO TRABALHO" in desc_n:
                if inicio is None:
                    inicio = t
                # eventos posteriores "INICIO DO TRABALHO" sao retomadas pos-pausa
            elif "INICIO DA PAUSA" in desc_n:
                if ini_pausa is None:
                    ini_pausa = t
            elif "FIM DA PAUSA" in desc_n:
                if fim_pausa is None:
                    fim_pausa = t
            elif "FINALIZACAO" in desc_n:
                fim = t
        m_trab = _RE_TRAB.search(bloco)
        m_pausa = _RE_PAUSA.search(bloco)
        trabalho = str_to_td(m_trab.group(1)) if m_trab else timedelta()
        pausa = str_to_td(m_pausa.group(1)) if m_pausa else timedelta()
        out.append({
            "data": d,
            "inicio": inicio,
            "ini_pausa": ini_pausa,
            "fim_pausa": fim_pausa,
            "fim": fim,
            "trabalho": trabalho,
            "pausa": pausa,
        })
    return out


# ---------- Excel Store ----------

class ExcelStore:
    def __init__(self, path):
        self.path = Path(path)
        # Workbook em cache na memoria: carregar o .xlsx (com estilos) custa
        # ~800ms; sem cache cada salvar/finalizar recarregava o arquivo varias
        # vezes. Mantemos uma unica copia viva e so recarregamos se o arquivo
        # mudar por fora (ex.: editado direto no Excel) — ver _get_wb.
        self._wb = None
        self._mtime = None
        if not self.path.exists():
            self._criar_inicial()

    def _criar_inicial(self):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_RESUMO
        self._setup_header(ws, HEADERS_RESUMO, [12, 18, 16, 14, 16, 18])
        ws2 = wb.create_sheet(SHEET_MARCACOES)
        self._setup_header(ws2, HEADERS_MARC, [14, 14, 40])
        wb.save(self.path)
        self._wb = wb
        self._mtime = self.path.stat().st_mtime

    def _setup_header(self, ws, headers, widths):
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = EX_HEADER_FONT
            c.fill = EX_HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = EX_BORDER
        ws.row_dimensions[1].height = 24
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    def _get_wb(self):
        """Retorna o workbook em cache. So recarrega do disco se o arquivo
        mudou por fora desde a ultima vez que o abrimos/salvamos (ex.: o
        usuario editou o .xlsx no Excel). Assim leituras e gravacoes
        sucessivas reaproveitam a mesma copia em memoria."""
        atual = self.path.stat().st_mtime if self.path.exists() else None
        if self._wb is None or atual != self._mtime:
            self._wb = load_workbook(self.path)
            self._mtime = atual
        return self._wb

    def _carregar(self):
        return self._get_wb()

    def _salvar(self, wb):
        wb.save(self.path)
        self._mtime = self.path.stat().st_mtime

    # ----- Marcacoes -----

    def listar_marcacoes(self):
        if not self.path.exists():
            return {}
        wb = self._get_wb()
        if SHEET_MARCACOES not in wb.sheetnames:
            return {}
        out = {}
        for row in wb[SHEET_MARCACOES].iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            d = parse_data(row[0])
            if not d:
                continue
            tipo = row[1] if len(row) > 1 else ""
            obs = row[2] if len(row) > 2 else ""
            out[d] = (tipo or "", obs or "")
        return out

    def set_marcacao(self, data, tipo, obs=""):
        wb = self._carregar()
        if SHEET_MARCACOES not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_MARCACOES)
            self._setup_header(ws, HEADERS_MARC, [14, 14, 40])
        else:
            ws = wb[SHEET_MARCACOES]
        row_idx = self._find_row_by_date(ws, data, col=1)
        if row_idx is None:
            row_idx = ws.max_row + 1
            if row_idx < 2:
                row_idx = 2
        ws.cell(row=row_idx, column=1, value=data.strftime("%d/%m/%Y"))
        ws.cell(row=row_idx, column=2, value=tipo)
        ws.cell(row=row_idx, column=3, value=obs)
        self._reordenar_marcacoes(ws)
        self._salvar(wb)

    def remover_marcacao(self, data):
        wb = self._carregar()
        if SHEET_MARCACOES not in wb.sheetnames:
            return
        ws = wb[SHEET_MARCACOES]
        row_idx = self._find_row_by_date(ws, data, col=1)
        if row_idx:
            ws.delete_rows(row_idx, 1)
            self._salvar(wb)

    def _reordenar_marcacoes(self, ws):
        rows = []
        for r in range(2, ws.max_row + 1):
            valores = [ws.cell(row=r, column=c).value for c in range(1, 4)]
            if valores[0]:
                rows.append(valores)
        rows.sort(key=lambda r: parse_data(r[0]) or date.min)
        # Apagar as linhas de dados de fato (delete_rows) em vez de so esvaziar:
        # o clear antigo percorria max_row+1, materializando uma celula nova a
        # cada save -> max_row crescia infinitamente (milhares de linhas-fantasma
        # estilizadas incharam o arquivo e deixaram o load lento).
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for i, vals in enumerate(rows, 2):
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.alignment = Alignment(horizontal="center" if c < 3 else "left", vertical="center")
                cell.border = EX_BORDER
                if i % 2 == 1:
                    cell.fill = EX_ALT_FILL

    # ----- Dias -----

    def _garantir_aba_mes(self, wb, ano, mes):
        nome = f"{ano:04d}-{mes:02d}"
        if nome in wb.sheetnames:
            return wb[nome]
        ws = wb.create_sheet(nome)
        self._setup_header(ws, HEADERS_DIA, [12, 10, 10, 10, 10, 10, 11, 11, 10, 10, 12, 28])
        return ws

    def _find_row_by_date(self, ws, data, col):
        data_str = data.strftime("%d/%m/%Y")
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v == data_str:
                return r
            d = parse_data(v)
            if d == data:
                return r
        return None

    def salvar_dia(self, data, inicio, ini_pausa, fim_pausa, fim, tipo="NORMAL", obs=""):
        """Persiste um dia. Recalcula totais e resumo."""
        wb = self._carregar()
        ws = self._garantir_aba_mes(wb, data.year, data.month)
        self._gravar_linha_dia(ws, data, inicio, ini_pausa, fim_pausa, fim, tipo, obs)
        # So a aba do mes afetado muda ao gravar um dia — reordenar TODAS as
        # abas de mes a cada save custava ~8s (era o gargalo do "Finalizar").
        # As demais ja estao ordenadas/estilizadas de gravacoes anteriores.
        self._reordenar_aba_mes(ws)
        self._atualizar_resumo(wb)
        self._salvar(wb)

    def salvar_dias_bulk(self, dias):
        wb = self._carregar()
        for d in dias:
            ws = self._garantir_aba_mes(wb, d["data"].year, d["data"].month)
            self._gravar_linha_dia(
                ws, d["data"],
                d.get("inicio"), d.get("ini_pausa"),
                d.get("fim_pausa"), d.get("fim"),
                d.get("tipo", "NORMAL"), d.get("obs", ""),
            )
        for name in list(wb.sheetnames):
            if re.match(r"^\d{4}-\d{2}$", name):
                self._reordenar_aba_mes(wb[name])
        self._atualizar_resumo(wb)
        self._salvar(wb)

    def excluir_dia(self, data):
        wb = self._carregar()
        nome = f"{data.year:04d}-{data.month:02d}"
        if nome not in wb.sheetnames:
            return
        ws = wb[nome]
        row_idx = self._find_row_by_date(ws, data, col=COL_DATA)
        if row_idx:
            ws.delete_rows(row_idx, 1)
            self._reordenar_aba_mes(ws)
            self._atualizar_resumo(wb)
            self._salvar(wb)

    def _gravar_linha_dia(self, ws, data, inicio, ini_pausa, fim_pausa, fim, tipo, obs):
        row_idx = self._find_row_by_date(ws, data, col=COL_DATA)
        if row_idx is None:
            row_idx = ws.max_row + 1
            if row_idx < 2:
                row_idx = 2
        if tipo == "NORMAL":
            trab, pausa = calcular_trab_pausa(inicio, ini_pausa, fim_pausa, fim)
            saldo = trab - META_DIARIA if inicio and fim else timedelta()
        else:
            trab, pausa = calcular_trab_pausa(inicio, ini_pausa, fim_pausa, fim)
            saldo = timedelta(0)
        valores = {
            COL_DATA: data.strftime("%d/%m/%Y"),
            COL_DIA: DIAS_SEMANA[data.weekday()],
            COL_INICIO: time_to_str(inicio),
            COL_INI_PAUSA: time_to_str(ini_pausa),
            COL_FIM_PAUSA: time_to_str(fim_pausa),
            COL_FIM: time_to_str(fim),
            COL_TRABALHO: td_to_str(trab) if tipo == "NORMAL" else "-",
            COL_PAUSA: td_to_str(pausa) if tipo == "NORMAL" else "-",
            COL_META: td_to_str(META_DIARIA) if tipo == "NORMAL" else "-",
            COL_SALDO: td_to_str(saldo) if tipo == "NORMAL" else "-",
            COL_TIPO: tipo,
            COL_OBS: obs,
        }
        for col, v in valores.items():
            ws.cell(row=row_idx, column=col, value=v)

    def _reordenar_aba_mes(self, ws):
        rows = []
        for r in range(2, ws.max_row + 1):
            valores = [ws.cell(row=r, column=c).value for c in range(1, len(HEADERS_DIA) + 1)]
            if valores[0] and str(valores[0]).strip() != "Total":
                rows.append(valores)
        rows.sort(key=lambda r: parse_data(r[0]) or date.min)
        # Apagar as linhas de fato (delete_rows) em vez de so esvaziar — ver
        # nota em _reordenar_marcacoes: o clear antigo inflava max_row a cada
        # save, gerando milhares de linhas-fantasma.
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        cores_tipo = {
            "FERIADO": "FFD8A8",
            "ATESTADO": "FFC9C9",
            "FOLGA": "D0EBFF",
        }
        for i, vals in enumerate(rows, 2):
            tipo = vals[COL_TIPO - 1] or "NORMAL"
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.alignment = Alignment(
                    horizontal="center" if c < COL_OBS else "left",
                    vertical="center",
                )
                cell.border = EX_BORDER
                if i % 2 == 1:
                    cell.fill = EX_ALT_FILL
            saldo_str = vals[COL_SALDO - 1]
            if tipo == "NORMAL" and saldo_str and saldo_str != "-":
                saldo_td = str_to_td(saldo_str)
                cell_saldo = ws.cell(row=i, column=COL_SALDO)
                if saldo_td.total_seconds() > 0:
                    cell_saldo.font = Font(color="2F9E44", bold=True)
                elif saldo_td.total_seconds() < 0:
                    cell_saldo.font = Font(color="E03131", bold=True)
            cor = cores_tipo.get(tipo)
            if cor:
                tipo_cell = ws.cell(row=i, column=COL_TIPO)
                tipo_cell.fill = PatternFill("solid", fgColor=cor)
                tipo_cell.font = Font(bold=True)

        normais = [r for r in rows if (r[COL_TIPO - 1] or "NORMAL") == "NORMAL"]
        neutros = [r for r in rows if (r[COL_TIPO - 1] or "NORMAL") in TIPOS_NEUTRALIZA_META]
        soma_trab_normal = sum((str_to_td(r[COL_TRABALHO - 1]) for r in normais), timedelta())
        soma_pausa = sum((str_to_td(r[COL_PAUSA - 1]) for r in normais), timedelta())
        soma_trab = soma_trab_normal + META_DIARIA * len(neutros)
        try:
            ano_aba, mes_aba = (int(x) for x in ws.title.split("-"))
            soma_meta = META_DIARIA * dias_uteis_no_mes(ano_aba, mes_aba)
        except (ValueError, AttributeError):
            soma_meta = META_DIARIA * len(normais)
        soma_saldo = soma_trab - soma_meta
        total_row = len(rows) + 2
        totais = {
            COL_DATA: "Total",
            COL_TRABALHO: td_to_str(soma_trab),
            COL_PAUSA: td_to_str(soma_pausa),
            COL_META: td_to_str(soma_meta),
            COL_SALDO: td_to_str(soma_saldo),
            COL_TIPO: f"{len(normais)} dias",
        }
        for c in range(1, len(HEADERS_DIA) + 1):
            v = totais.get(c, "")
            cell = ws.cell(row=total_row, column=c, value=v)
            cell.font = Font(bold=True)
            cell.fill = EX_TOTAL_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = EX_TOTAL_BORDER
        if soma_saldo.total_seconds() > 0:
            ws.cell(row=total_row, column=COL_SALDO).font = Font(color="2F9E44", bold=True)
        elif soma_saldo.total_seconds() < 0:
            ws.cell(row=total_row, column=COL_SALDO).font = Font(color="E03131", bold=True)

    def listar_meses(self):
        if not self.path.exists():
            return []
        wb = self._get_wb()
        out = []
        for name in wb.sheetnames:
            if re.match(r"^\d{4}-\d{2}$", name):
                a, m = name.split("-")
                out.append((int(a), int(m)))
        return sorted(out)

    def listar_dias_do_mes(self, ano, mes):
        if not self.path.exists():
            return []
        wb = self._get_wb()
        nome = f"{ano:04d}-{mes:02d}"
        if nome not in wb.sheetnames:
            return []
        out = []
        for row in wb[nome].iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            if str(row[0]).strip() == "Total":
                continue
            d = parse_data(row[0])
            if not d:
                continue
            inicio = combinar(d, parse_hhmmss(row[COL_INICIO - 1]))
            ini_pausa = combinar(d, parse_hhmmss(row[COL_INI_PAUSA - 1]))
            fim_pausa = combinar(d, parse_hhmmss(row[COL_FIM_PAUSA - 1]))
            fim = combinar(d, parse_hhmmss(row[COL_FIM - 1]))
            out.append({
                "data": d,
                "dia_semana": row[COL_DIA - 1] or DIAS_SEMANA[d.weekday()],
                "inicio": inicio,
                "ini_pausa": ini_pausa,
                "fim_pausa": fim_pausa,
                "fim": fim,
                "trabalho": str_to_td(row[COL_TRABALHO - 1]),
                "pausa": str_to_td(row[COL_PAUSA - 1]),
                "saldo": str_to_td(row[COL_SALDO - 1]) if row[COL_SALDO - 1] and row[COL_SALDO - 1] != "-" else timedelta(0),
                "tipo": row[COL_TIPO - 1] or "NORMAL",
                "obs": row[COL_OBS - 1] if len(row) >= COL_OBS and row[COL_OBS - 1] else "",
            })
        return out

    def get_dia(self, data):
        for d in self.listar_dias_do_mes(data.year, data.month):
            if d["data"] == data:
                return d
        return None

    def _atualizar_resumo(self, wb):
        if SHEET_RESUMO not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_RESUMO, 0)
            self._setup_header(ws, HEADERS_RESUMO, [12, 18, 16, 14, 16, 18])
        else:
            ws = wb[SHEET_RESUMO]
            # Apagar as linhas de fato (delete_rows) em vez de so esvaziar — ver
            # nota em _reordenar_marcacoes.
            if ws.max_row >= 2:
                ws.delete_rows(2, ws.max_row - 1)
        meses = []
        for name in wb.sheetnames:
            if re.match(r"^\d{4}-\d{2}$", name):
                a, m = name.split("-")
                meses.append((int(a), int(m), name))
        meses.sort()
        acumulado = timedelta()
        for i, (a, m, name) in enumerate(meses, 2):
            ws_m = wb[name]
            soma_trab = timedelta()
            dias = 0
            for row in ws_m.iter_rows(min_row=2, values_only=True):
                if not row or not row[0] or str(row[0]).strip() == "Total":
                    continue
                tipo = row[COL_TIPO - 1] or "NORMAL"
                if tipo == "NORMAL":
                    soma_trab += str_to_td(row[COL_TRABALHO - 1])
                    dias += 1
                elif tipo in TIPOS_NEUTRALIZA_META:
                    soma_trab += META_DIARIA
            soma_meta = META_DIARIA * dias_uteis_no_mes(a, m)
            saldo_mes = soma_trab - soma_meta
            acumulado += saldo_mes
            valores = [
                f"{a:04d}-{m:02d}", str(dias),
                td_to_str(soma_trab), td_to_str(soma_meta),
                td_to_str(saldo_mes), td_to_str(acumulado),
            ]
            for col, v in enumerate(valores, 1):
                cell = ws.cell(row=i, column=col, value=v)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = EX_BORDER
                if i % 2 == 1:
                    cell.fill = EX_ALT_FILL
            for col_color, td_val in [(5, saldo_mes), (6, acumulado)]:
                cell = ws.cell(row=i, column=col_color)
                if td_val.total_seconds() > 0:
                    cell.font = Font(color="2F9E44", bold=True)
                elif td_val.total_seconds() < 0:
                    cell.font = Font(color="E03131", bold=True)


# ---------- Cálculo de saldo ----------

def calcular_saldos(store):
    hoje = date.today()
    marcacoes = store.listar_marcacoes()
    meses = store.listar_meses()
    trabalhados = {}
    for a, m in meses:
        for d in store.listar_dias_do_mes(a, m):
            trabalhados[d["data"]] = d

    if not trabalhados and not marcacoes:
        return timedelta(), timedelta(), []

    primeiro = min(list(trabalhados.keys()) + list(marcacoes.keys()))
    saldo_geral = timedelta()
    saldo_mes = timedelta()
    entradas_mes = []

    cur = primeiro
    while cur <= hoje:
        no_mes_atual = (cur.year == hoje.year and cur.month == hoje.month)
        if cur.weekday() < 5:
            if cur in marcacoes:
                tipo, obs = marcacoes[cur]
                entry = {
                    "data": cur, "dia_semana": DIAS_SEMANA[cur.weekday()],
                    "trabalho": timedelta(), "pausa": timedelta(),
                    "saldo": timedelta(), "tipo": tipo or "FERIADO", "obs": obs,
                }
                if cur in trabalhados:
                    t = trabalhados[cur]
                    entry["trabalho"] = t["trabalho"]
                    entry["pausa"] = t["pausa"]
            elif cur in trabalhados:
                t = trabalhados[cur]
                tipo = t["tipo"]
                if tipo == "NORMAL":
                    contrib = t["trabalho"] - META_DIARIA if t["inicio"] and t["fim"] else timedelta()
                else:
                    contrib = timedelta()
                entry = {
                    "data": cur, "dia_semana": t["dia_semana"],
                    "trabalho": t["trabalho"], "pausa": t["pausa"],
                    "saldo": contrib, "tipo": tipo, "obs": t.get("obs", ""),
                }
            else:
                entry = {
                    "data": cur, "dia_semana": DIAS_SEMANA[cur.weekday()],
                    "trabalho": timedelta(), "pausa": timedelta(),
                    "saldo": -META_DIARIA, "tipo": "FALTA",
                    "obs": "Dia util sem registro",
                }
            saldo_geral += entry["saldo"]
            if no_mes_atual:
                saldo_mes += entry["saldo"]
                entradas_mes.append(entry)
        else:
            if cur in trabalhados:
                t = trabalhados[cur]
                if t["tipo"] == "NORMAL" and t["inicio"] and t["fim"]:
                    contrib = t["trabalho"]
                    saldo_geral += contrib
                    entry = {
                        "data": cur, "dia_semana": t["dia_semana"],
                        "trabalho": t["trabalho"], "pausa": t["pausa"],
                        "saldo": contrib, "tipo": "EXTRA",
                        "obs": "Fim de semana",
                    }
                    if no_mes_atual:
                        saldo_mes += contrib
                        entradas_mes.append(entry)
        cur += timedelta(days=1)

    return saldo_mes, saldo_geral, entradas_mes


# ---------- Widgets custom ----------

class HoverButton(tk.Canvas):
    """Botao com cantos arredondados desenhado via Canvas + Pillow.
    Mantem a API do antigo tk.Button-based HoverButton (config bg/fg/text/state,
    set_palette, grid/pack/bind) para nao quebrar callsites."""

    def __init__(self, master, bg=COR_PRIMARY, hover_bg=COR_PRIMARY_HOVER,
                 fg="white", font=("Segoe UI", 10, "bold"),
                 padx=18, pady=10, radius=12, text="", command=None,
                 parent_bg=None, **kwargs):
        self._bg_normal = bg
        self._bg_hover = hover_bg
        self._fg_normal = fg
        self._text = text
        self._font = font
        self._padx = padx
        self._pady = pady
        self._radius = radius
        self._command = command
        self._state = "normal"
        self._parent_bg = parent_bg or self._inferir_parent_bg(master)
        self._photo = None

        import tkinter.font as tkfont
        tkf = tkfont.Font(font=font)
        tw = tkf.measure(text or " ")
        th = tkf.metrics("linespace")
        w = tw + 2 * padx
        h = th + 2 * pady

        kwargs.pop("relief", None)
        kwargs.pop("borderwidth", None)
        kwargs.setdefault("cursor", "hand2")
        kwargs.setdefault("highlightthickness", 0)
        kwargs.setdefault("bd", 0)
        super().__init__(master, width=w, height=h, bg=self._parent_bg, **kwargs)

        self._redesenhar(self._bg_normal)
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<Button-1>", self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)

    @staticmethod
    def _inferir_parent_bg(master):
        try:
            return master.cget("bg")
        except tk.TclError:
            return COR_BG

    def _redesenhar(self, cor_fundo):
        w = int(self.cget("width"))
        h = int(self.cget("height"))
        if w < 2 or h < 2:
            return
        if PIL_OK:
            base_rgb = _hex_para_rgb(self._parent_bg)
            img = Image.new("RGBA", (w, h), base_rgb + (255,))
            d = ImageDraw.Draw(img)
            cor = cor_fundo if self._state == "normal" else COR_DISABLED
            d.rounded_rectangle(
                (0, 0, w - 1, h - 1), radius=self._radius,
                fill=_hex_para_rgb(cor) + (255,),
            )
            self._photo = ImageTk.PhotoImage(img)
            self.delete("all")
            self.create_image(0, 0, image=self._photo, anchor="nw")
        else:
            self.delete("all")
            cor = cor_fundo if self._state == "normal" else COR_DISABLED
            self.create_rectangle(0, 0, w, h, fill=cor, outline="")
        fg = self._fg_normal if self._state == "normal" else "white"
        self.create_text(
            w / 2, h / 2, text=self._text, fill=fg, font=self._font,
        )

    def _on_enter(self, _ev):
        if self._state == "normal":
            self._redesenhar(self._bg_hover)

    def _on_leave(self, _ev):
        if self._state == "normal":
            self._redesenhar(self._bg_normal)

    def _on_press(self, _ev):
        if self._state == "normal":
            self._redesenhar(self._bg_hover)

    def _on_release(self, ev):
        if self._state != "normal":
            return
        self._redesenhar(self._bg_normal)
        if 0 <= ev.x <= int(self.cget("width")) and 0 <= ev.y <= int(self.cget("height")):
            if self._command:
                self._command()

    def config(self, cnf=None, **kwargs):
        if cnf:
            kwargs = {**cnf, **kwargs}
        precisa_redraw = False
        if "text" in kwargs:
            self._text = kwargs.pop("text")
            precisa_redraw = True
        if "fg" in kwargs:
            self._fg_normal = kwargs.pop("fg")
            precisa_redraw = True
        if "bg" in kwargs:
            self._bg_normal = kwargs.pop("bg")
            precisa_redraw = True
        if "state" in kwargs:
            self._state = kwargs.pop("state")
            super().configure(cursor="hand2" if self._state == "normal" else "arrow")
            precisa_redraw = True
        if "command" in kwargs:
            self._command = kwargs.pop("command")
        if kwargs:
            super().configure(**kwargs)
        if precisa_redraw:
            self._redesenhar(self._bg_normal)

    configure = config

    def __getitem__(self, key):
        if key == "state":
            return self._state
        return super().__getitem__(key)

    def set_palette(self, bg, hover_bg):
        self._bg_normal = bg
        self._bg_hover = hover_bg
        self._redesenhar(self._bg_normal)


class RoundedCard(tk.Frame):
    """Frame com fundo arredondado (e sombra suave opcional) desenhado via Pillow.
    Use `card.content` para colocar widgets dentro (igual a um tk.Frame normal)."""

    def __init__(self, master, bg=None, parent_bg=None,
                 radius=14, shadow=True, padding=0, border=True, **kwargs):
        bg = bg or COR_CARD
        parent_bg = parent_bg or self._inferir_parent_bg(master)
        kwargs.pop("bg", None)
        super().__init__(master, bg=parent_bg, **kwargs)
        self._bg = bg
        self._parent_bg = parent_bg
        self._radius = radius
        self._shadow = shadow
        self._border = border
        self._photo = None
        self._last_size = (0, 0)

        self._canvas = tk.Canvas(
            self, bg=parent_bg, highlightthickness=0, bd=0,
        )
        self._canvas.place(x=0, y=0, relwidth=1, relheight=1)
        self.content = tk.Frame(self, bg=bg)
        m = padding + (4 if shadow else 0)
        self.content.place(
            x=m, y=m,
            relwidth=1, relheight=1,
            width=-2 * m, height=-2 * m,
        )
        self.bind("<Configure>", self._on_resize)

    @staticmethod
    def _inferir_parent_bg(master):
        try:
            return master.cget("bg")
        except tk.TclError:
            return COR_BG

    def _on_resize(self, _ev):
        self._redesenhar()

    def _redesenhar(self):
        w = self.winfo_width()
        h = self.winfo_height()
        if w < 4 or h < 4 or (w, h) == self._last_size:
            return
        self._last_size = (w, h)
        if not PIL_OK:
            self._canvas.delete("all")
            self._canvas.create_rectangle(0, 0, w, h, fill=self._bg, outline="")
            return
        base = Image.new("RGBA", (w, h), _hex_para_rgb(self._parent_bg) + (255,))
        if self._shadow:
            sombra = Image.new("RGBA", (w, h), (0, 0, 0, 0))
            ds = ImageDraw.Draw(sombra)
            offset = 4
            ds.rounded_rectangle(
                (offset, offset + 2, w - offset, h - offset + 1),
                radius=self._radius, fill=(15, 23, 42, 38),
            )
            sombra = sombra.filter(ImageFilter.GaussianBlur(6))
            base = Image.alpha_composite(base, sombra)
        topo = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        dt = ImageDraw.Draw(topo)
        margem = 4 if self._shadow else 0
        if self._border:
            dt.rounded_rectangle(
                (margem, margem, w - margem - 1, h - margem - 1),
                radius=self._radius, fill=_hex_para_rgb(COR_BORDER) + (255,),
            )
            dt.rounded_rectangle(
                (margem + 1, margem + 1, w - margem - 2, h - margem - 2),
                radius=max(0, self._radius - 1),
                fill=_hex_para_rgb(self._bg) + (255,),
            )
        else:
            dt.rounded_rectangle(
                (margem, margem, w - margem - 1, h - margem - 1),
                radius=self._radius, fill=_hex_para_rgb(self._bg) + (255,),
            )
        base = Image.alpha_composite(base, topo)
        self._photo = ImageTk.PhotoImage(base)
        self._canvas.delete("all")
        self._canvas.create_image(0, 0, image=self._photo, anchor="nw")


def mostrar_toast(root, mensagem, cor=COR_SUCCESS, duracao=2200):
    """Janelinha temporaria no canto inferior direito (no-op visual quando recolhido)."""
    try:
        toast = tk.Toplevel(root)
        toast.overrideredirect(True)
        toast.attributes("-topmost", True)
        try:
            toast.attributes("-alpha", 0.96)
        except tk.TclError:
            pass
        # tamanho proporcional ao texto
        w = max(220, min(420, len(mensagem) * 9 + 40))
        h = 54
        root.update_idletasks()
        rx = root.winfo_rootx()
        ry = root.winfo_rooty()
        rw = root.winfo_width()
        rh = root.winfo_height()
        x = rx + rw - w - 24
        y = ry + rh - h - 24
        toast.geometry(f"{w}x{h}+{x}+{y}")
        wrap = tk.Frame(toast, bg=cor, bd=0, highlightthickness=0)
        wrap.pack(fill="both", expand=True)
        tk.Label(
            wrap, text=mensagem, bg=cor, fg="white",
            font=("Segoe UI", 10, "bold"), padx=14, pady=10,
        ).pack(expand=True)
        toast.after(duracao, toast.destroy)
    except tk.TclError:
        pass


# ---------- UI ----------

class FolhaPontoApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Folha de Ponto")
        self.root.configure(bg=COR_BG)
        self.root.geometry("860x720")
        self.root.minsize(780, 660)
        self.root.attributes("-topmost", True)
        self.root.after(800, lambda: self.root.attributes("-topmost", False))
        self._aplicar_icone()

        self.store = None
        if OPENPYXL_OK:
            try:
                self.store = ExcelStore(EXCEL_PATH)
                if not any(self.store.listar_meses()):
                    self._migrar_txts_iniciais()
            except PermissionError:
                messagebox.showwarning(
                    "Excel em uso",
                    f"Nao consegui acessar {EXCEL_PATH.name}. "
                    "Feche o arquivo no Excel e abra o app de novo."
                )

        # Estado "ao vivo" do dia atual (cronometro)
        self.hoje_dia = self._dia_vazio()
        if self.store:
            try:
                d_salvo = self.store.get_dia(date.today())
                if d_salvo:
                    self.hoje_dia = {
                        "inicio": d_salvo["inicio"],
                        "ini_pausa": d_salvo["ini_pausa"],
                        "fim_pausa": d_salvo["fim_pausa"],
                        "fim": d_salvo["fim"],
                        "tipo": d_salvo["tipo"],
                        "obs": d_salvo["obs"],
                    }
            except Exception as e:
                print(f"Erro restaurando estado de hoje: {e}")
        self.hoje_eventos = []
        # Dia atualmente visualizado/editado na aba Dia
        self.data_vis = date.today()
        # Espelho dos campos editaveis (sempre os 4 datetimes ou None)
        self.dia_vis = self._dia_vazio()

        self._setup_styles()
        self._setup_ui()
        if not OPENPYXL_OK:
            messagebox.showerror(
                "openpyxl ausente",
                "Modulo openpyxl nao esta instalado.\n\n"
                "Instale com:\n  python -m pip install openpyxl"
            )
        self._carregar_dia_vis()
        self.atualizar_tempo()
        self.atualizar_visoes()
        self.root.after(AUTOSAVE_INTERVALO_MS, self._autosave_hoje)

    def _aplicar_icone(self):
        # Windows: forca AppUserModelID antes pra que a barra de tarefas
        # use o icone do app em vez do icone do interpretador Python.
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                APP_USER_MODEL_ID
            )
        except Exception:
            pass
        if ICONE_PATH.exists():
            try:
                self.root.iconbitmap(default=str(ICONE_PATH))
            except Exception as e:
                print(f"Falha ao carregar icone: {e}")

    def _dia_vazio(self):
        return {
            "inicio": None, "ini_pausa": None, "fim_pausa": None, "fim": None,
            "tipo": "NORMAL", "obs": "",
        }

    def _migrar_txts_iniciais(self):
        arquivos = sorted(DATA_DIR.glob("ponto_*.txt"))
        if not arquivos:
            return
        dias = []
        for arq in arquivos:
            dias.extend(parse_blocos_txt(arq))
        if not dias:
            return
        mapa = {d["data"]: d for d in dias}
        self.store.salvar_dias_bulk(sorted(mapa.values(), key=lambda x: x["data"]))
        messagebox.showinfo(
            "Migracao concluida",
            f"Importados {len(mapa)} dias dos arquivos .txt para {EXCEL_PATH.name}."
        )

    # ----- estilos -----

    def _setup_styles(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(
            "TNotebook", background=COR_BG, borderwidth=0,
            tabmargins=[12, 6, 12, 0],
        )
        style.configure(
            "TNotebook.Tab", padding=[26, 12],
            font=("Segoe UI", 10, "bold"),
            background=COR_BG, foreground=COR_MUTED, borderwidth=0,
            focuscolor=COR_BG,
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", COR_CARD), ("active", COR_HOVER)],
            foreground=[("selected", COR_PRIMARY), ("active", COR_TEXT)],
            expand=[("selected", [1, 1, 1, 0])],
        )
        style.configure("TFrame", background=COR_BG)
        style.configure("Card.TFrame", background=COR_CARD)
        style.configure("TLabel", background=COR_BG, foreground=COR_TEXT, font=("Segoe UI", 10))
        style.configure("Card.TLabel", background=COR_CARD, foreground=COR_TEXT, font=("Segoe UI", 10))
        style.configure("Title.TLabel", background=COR_CARD, foreground=COR_TEXT, font=("Segoe UI", 13, "bold"))
        style.configure("Muted.TLabel", background=COR_CARD, foreground=COR_MUTED, font=("Segoe UI", 9))
        style.configure("Big.TLabel", background=COR_CARD, foreground=COR_PRIMARY, font=("Segoe UI", 40, "bold"))
        style.configure(
            "TProgressbar", background=COR_PRIMARY,
            troughcolor=COR_BORDER, borderwidth=0, thickness=10,
        )
        style.configure(
            "Day.Horizontal.TProgressbar",
            background=COR_PRIMARY, troughcolor=COR_BORDER,
            borderwidth=0, thickness=10, lightcolor=COR_PRIMARY,
            darkcolor=COR_PRIMARY,
        )
        style.configure(
            "Treeview", background=COR_CARD, foreground=COR_TEXT,
            fieldbackground=COR_CARD, font=("Segoe UI", 10),
            rowheight=36, borderwidth=0,
        )
        style.configure(
            "Treeview.Heading", background=COR_HOVER, foreground=COR_TEXT_SOFT,
            font=("Segoe UI", 9, "bold"), borderwidth=0, padding=10,
            relief="flat",
        )
        style.map(
            "Treeview.Heading",
            background=[("active", COR_BORDER)],
        )
        style.map(
            "Treeview",
            background=[("selected", COR_PRIMARY)],
            foreground=[("selected", "white")],
        )
        style.configure(
            "TCombobox",
            fieldbackground=COR_CARD, background=COR_CARD,
            foreground=COR_TEXT, padding=4,
        )

    # ----- estrutura -----

    def _setup_ui(self):
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=14, pady=14)
        self.tab_dia = ttk.Frame(self.notebook, style="Card.TFrame")
        self.tab_mes = ttk.Frame(self.notebook, style="Card.TFrame")
        self.tab_geral = ttk.Frame(self.notebook, style="Card.TFrame")
        self.tab_marc = ttk.Frame(self.notebook, style="Card.TFrame")
        self.notebook.add(self.tab_dia, text="  📅  Dia  ")
        self.notebook.add(self.tab_mes, text="  📆  Mes  ")
        self.notebook.add(self.tab_geral, text="  🏦  Banco de horas  ")
        self.notebook.add(self.tab_marc, text="  🏷  Marcacoes  ")
        self._build_tab_dia()
        self._build_tab_mes()
        self._build_tab_geral()
        self._build_tab_marc()
        self.notebook.bind("<<NotebookTabChanged>>", lambda e: self.atualizar_visoes())
        # Atalhos de navegacao
        self.root.bind("<Left>", self._kb_nav_anterior)
        self.root.bind("<Right>", self._kb_nav_proximo)
        self.root.bind("<Home>", lambda e: self._kb_ir_hoje())
        self.root.bind("<Control-s>", lambda e: self._kb_salvar())

    def _kb_focus_em_entry(self):
        w = self.root.focus_get()
        return isinstance(w, (tk.Entry, ttk.Combobox))

    def _kb_nav_anterior(self, ev):
        if self._kb_focus_em_entry():
            return
        if self.notebook.index(self.notebook.select()) == 0:
            self.navegar_dia(-1)

    def _kb_nav_proximo(self, ev):
        if self._kb_focus_em_entry():
            return
        if self.notebook.index(self.notebook.select()) == 0:
            self.navegar_dia(1)

    def _kb_ir_hoje(self):
        if self.notebook.index(self.notebook.select()) == 0:
            self.ir_para_hoje()

    def _kb_salvar(self):
        if self.notebook.index(self.notebook.select()) == 0:
            self.salvar_dia_atual()

    # ----- aba Dia -----

    def _build_tab_dia(self):
        f = self.tab_dia

        # ---- Header com carrossel ----
        # place() em cada bloco: prev colado na borda esquerda, right_group
        # colado na direita, e a data no centro geometrico do hdr
        # (relx=0.5, anchor="center"). Como hdr ocupa toda a largura util da
        # aba (pack fill="x"), o centro do hdr coincide com o centro
        # horizontal da janela.
        hdr = tk.Frame(f, bg=COR_CARD, height=64)
        hdr.pack(fill="x", padx=24, pady=(18, 4))
        hdr.pack_propagate(False)

        self.btn_prev_dia = HoverButton(
            hdr, text="◀", bg=COR_CARD, hover_bg=COR_HOVER, fg=COR_PRIMARY,
            font=("Segoe UI", 16, "bold"), padx=14, pady=4,
            command=lambda: self.navegar_dia(-1),
        )
        self.btn_prev_dia.place(relx=0, rely=0.5, anchor="w")

        centro = tk.Frame(hdr, bg=COR_CARD)
        centro.place(relx=0.5, rely=0.5, anchor="center")
        self.lbl_data_principal = tk.Label(
            centro, text="", bg=COR_CARD, fg=COR_TEXT,
            font=("Segoe UI Semibold", 15),
        )
        self.lbl_data_principal.pack()
        self.lbl_data_sub = tk.Label(
            centro, text="", bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 9),
        )
        self.lbl_data_sub.pack(pady=(2, 0))

        right_group = tk.Frame(hdr, bg=COR_CARD)
        right_group.place(relx=1, rely=0.5, anchor="e")
        self.btn_hoje = HoverButton(
            right_group, text="● Hoje", bg=COR_PRIMARY_LIGHT, hover_bg="#c7d2fe",
            fg=COR_PRIMARY_TEXT, font=("Segoe UI", 9, "bold"),
            padx=12, pady=6, command=self.ir_para_hoje,
        )
        self.btn_hoje.pack(side="left", padx=(0, 10))
        self.btn_next_dia = HoverButton(
            right_group, text="▶", bg=COR_CARD, hover_bg=COR_HOVER, fg=COR_PRIMARY,
            font=("Segoe UI", 16, "bold"), padx=14, pady=4,
            command=lambda: self.navegar_dia(1),
        )
        self.btn_next_dia.pack(side="left")

        # ---- Timer grande + progresso ----
        timer_wrap = tk.Frame(f, bg=COR_CARD)
        timer_wrap.pack(pady=(10, 0))
        self.lbl_timer_big = tk.Label(
            timer_wrap, text="00:00:00", bg=COR_CARD, fg=COR_PRIMARY,
            font=("Consolas", 42, "bold"),
        )
        self.lbl_timer_big.pack()
        self.lbl_timer_sub = tk.Label(
            timer_wrap, text="Trabalhado hoje", bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 9),
        )
        self.lbl_timer_sub.pack()

        prog_frame = tk.Frame(f, bg=COR_CARD)
        prog_frame.pack(fill="x", padx=80, pady=(12, 6))
        lbl_row = tk.Frame(prog_frame, bg=COR_CARD)
        lbl_row.pack(fill="x")
        self.lbl_pct = tk.Label(
            lbl_row, text="0% da meta diaria", bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 9),
        )
        self.lbl_pct.pack(side="left")
        self.lbl_meta_alvo = tk.Label(
            lbl_row, text=f"meta {td_to_str(META_DIARIA)}", bg=COR_CARD,
            fg=COR_MUTED, font=("Segoe UI", 9),
        )
        self.lbl_meta_alvo.pack(side="right")
        self.progress = ttk.Progressbar(
            prog_frame, length=400, mode="determinate", maximum=100,
            style="Day.Horizontal.TProgressbar",
        )
        self.progress.pack(fill="x", pady=(4, 0))

        # ---- Cards de status (em pausa, saldo, status) ----
        info = tk.Frame(f, bg=COR_CARD)
        info.pack(pady=(12, 4))
        self.card_pausa_dia = self._make_status_card(info, "EM PAUSA", "00:00:00", COR_WARNING)
        self.card_pausa_dia["frame"].grid(row=0, column=0, padx=8)
        self.card_saldo_dia = self._make_status_card(info, "SALDO DO DIA", "-08:00:00", COR_DANGER)
        self.card_saldo_dia["frame"].grid(row=0, column=1, padx=8)
        self.card_status_dia = self._make_status_pill(info, "STATUS", "Parado", COR_MUTED)
        self.card_status_dia["frame"].grid(row=0, column=2, padx=8)

        # ---- Timeline horizontal dos 4 horarios ----
        time_wrap = tk.Frame(f, bg=COR_CARD)
        time_wrap.pack(fill="x", padx=24, pady=(16, 4))

        self.var_inicio = tk.StringVar()
        self.var_ini_pausa = tk.StringVar()
        self.var_fim_pausa = tk.StringVar()
        self.var_fim = tk.StringVar()
        self.var_tipo = tk.StringVar(value="NORMAL")
        self.var_obs = tk.StringVar()

        timeline = tk.Frame(time_wrap, bg=COR_CARD)
        timeline.pack(fill="x")
        self._make_time_card(timeline, "▶  Inicio", self.var_inicio, "inicio", 0, COR_SUCCESS)
        self._make_arrow(timeline, 1)
        self._make_time_card(timeline, "⏸  Ini pausa", self.var_ini_pausa, "ini_pausa", 2, COR_WARNING)
        self._make_arrow(timeline, 3)
        self._make_time_card(timeline, "▶  Fim pausa", self.var_fim_pausa, "fim_pausa", 4, COR_WARNING)
        self._make_arrow(timeline, 5)
        self._make_time_card(timeline, "⏹  Fim", self.var_fim, "fim", 6, COR_DANGER)
        for i in (0, 2, 4, 6):
            timeline.columnconfigure(i, weight=1, uniform="time")

        # ---- Tipo + Obs ----
        meta = tk.Frame(f, bg=COR_CARD)
        meta.pack(fill="x", padx=24, pady=(12, 4))
        tk.Label(meta, text="Tipo", bg=COR_CARD, fg=COR_MUTED,
                 font=("Segoe UI", 9, "bold")).grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.cb_tipo = ttk.Combobox(
            meta, textvariable=self.var_tipo, values=TIPOS_DIA,
            state="readonly", width=10, font=("Segoe UI", 10),
        )
        self.cb_tipo.grid(row=0, column=1, sticky="w")
        self.cb_tipo.bind("<<ComboboxSelected>>", lambda e: self._on_tipo_change())

        tk.Label(meta, text="Observacao", bg=COR_CARD, fg=COR_MUTED,
                 font=("Segoe UI", 9, "bold")).grid(row=0, column=2, sticky="w", padx=(24, 8))
        self.entry_obs = tk.Entry(
            meta, textvariable=self.var_obs, font=("Segoe UI", 10),
            relief="flat", bd=0, highlightthickness=1,
            highlightbackground=COR_BORDER, highlightcolor=COR_PRIMARY,
        )
        self.entry_obs.grid(row=0, column=3, sticky="ew", ipady=6)
        meta.columnconfigure(3, weight=1)

        # ---- Botoes principais (live) ----
        btns = tk.Frame(f, bg=COR_CARD)
        btns.pack(pady=(16, 4))
        self.btn_play = HoverButton(
            btns, text="▶  Play / Retomar",
            bg=COR_SUCCESS, hover_bg=COR_SUCCESS_HOVER,
            padx=20, pady=11, command=self.play,
        )
        self.btn_play.grid(row=0, column=0, padx=4)
        self.btn_pause = HoverButton(
            btns, text="⏸  Pause",
            bg=COR_WARNING, hover_bg=COR_WARNING_HOVER,
            padx=20, pady=11, command=self.pause,
        )
        self.btn_pause.grid(row=0, column=1, padx=4)
        self.btn_fim_dia = HoverButton(
            btns, text="⏹  Finalizar dia",
            bg=COR_DANGER, hover_bg=COR_DANGER_HOVER,
            padx=20, pady=11, command=self.finalizar_dia,
        )
        self.btn_fim_dia.grid(row=0, column=2, padx=4)

        # ---- Botoes secundarios (salvar/excluir) ----
        btns2 = tk.Frame(f, bg=COR_CARD)
        btns2.pack(pady=(2, 8))
        self.btn_salvar = HoverButton(
            btns2, text="💾  Salvar",
            bg=COR_PRIMARY, hover_bg=COR_PRIMARY_HOVER,
            padx=18, pady=9, command=self.salvar_dia_atual,
        )
        self.btn_salvar.grid(row=0, column=0, padx=4)
        self.btn_excluir = HoverButton(
            btns2, text="🗑  Excluir dia",
            bg=COR_CARD, hover_bg=COR_DANGER_LIGHT,
            fg=COR_DANGER, font=("Segoe UI", 10, "bold"),
            padx=18, pady=9, command=self.excluir_dia_atual,
            highlightthickness=1, highlightbackground=COR_BORDER,
        )
        self.btn_excluir.grid(row=0, column=1, padx=4)

        tk.Label(
            f, text="←/→ navegar  •  Home volta pra hoje  •  Ctrl+S salva",
            bg=COR_CARD, fg=COR_MUTED, font=("Segoe UI", 8),
        ).pack(side="bottom", pady=(0, 10))

    def _make_time_card(self, parent, label, var, key, col, cor):
        card = tk.Frame(parent, bg=COR_CARD)
        card.grid(row=0, column=col, padx=4, sticky="ew")
        tk.Label(
            card, text=label, bg=COR_CARD, fg=cor,
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w", padx=2)
        e = tk.Entry(
            card, textvariable=var, font=("Consolas", 13),
            relief="flat", bd=0, justify="center",
            highlightthickness=2, highlightbackground=COR_BORDER,
            highlightcolor=cor,
        )
        e.pack(fill="x", ipady=8, pady=(2, 0))
        e.bind("<FocusOut>", lambda ev: self._commit_campo(key))
        e.bind("<Return>", lambda ev: (self._commit_campo(key), self.root.focus_set()))
        return e

    def _make_arrow(self, parent, col):
        arr = tk.Label(
            parent, text="→", bg=COR_CARD, fg=COR_BORDER_FORTE,
            font=("Segoe UI", 16, "bold"),
        )
        arr.grid(row=0, column=col, sticky="s", pady=(18, 0))

    def _make_status_card(self, parent, label, valor, cor_valor):
        try:
            parent_bg = parent.cget("bg")
        except tk.TclError:
            parent_bg = COR_CARD
        cf = RoundedCard(parent, bg=COR_CARD, parent_bg=parent_bg, radius=12)
        cf.configure(width=160, height=82)
        inner = tk.Frame(cf.content, bg=COR_CARD, padx=22, pady=12)
        inner.pack()
        val = tk.Label(
            inner, text=valor, bg=COR_CARD, fg=cor_valor,
            font=("Consolas", 16, "bold"),
        )
        val.pack()
        lbl = tk.Label(
            inner, text=label, bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 8, "bold"),
        )
        lbl.pack(pady=(2, 0))
        return {"frame": cf, "value": val, "label": lbl}

    def _make_status_pill(self, parent, label, valor, cor_valor):
        try:
            parent_bg = parent.cget("bg")
        except tk.TclError:
            parent_bg = COR_CARD
        cf = RoundedCard(parent, bg=COR_CARD, parent_bg=parent_bg, radius=12)
        cf.configure(width=160, height=82)
        inner = tk.Frame(cf.content, bg=COR_CARD, padx=22, pady=10)
        inner.pack()
        # Pill colorida com valor
        pill = tk.Frame(inner, bg=cor_valor)
        pill.pack()
        val = tk.Label(
            pill, text=valor, bg=cor_valor, fg="white",
            font=("Segoe UI", 11, "bold"), padx=14, pady=5,
        )
        val.pack()
        lbl = tk.Label(
            inner, text=label, bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 8, "bold"),
        )
        lbl.pack(pady=(4, 0))
        return {"frame": cf, "value": val, "label": lbl, "pill": pill}

    # ----- aba Mes atual -----

    def _build_tab_mes(self):
        f = self.tab_mes
        hoje = date.today()
        hdr = tk.Frame(f, bg=COR_CARD)
        hdr.pack(fill="x", padx=16, pady=(16, 6))
        ttk.Label(
            hdr, text=f"{MESES_PT[hoje.month - 1]} / {hoje.year}", style="Title.TLabel"
        ).pack(side="left")

        sumf = tk.Frame(f, bg=COR_CARD)
        sumf.pack(fill="x", padx=16, pady=6)
        self.card_trab_mes = self._mini_card(sumf, "Trabalhado", "00:00:00", COR_PRIMARY)
        self.card_meta_mes = self._mini_card(sumf, "Meta", "00:00:00", COR_MUTED)
        self.card_saldo_mes = self._mini_card(sumf, "Saldo do mes", "00:00:00", COR_SUCCESS)
        self.card_dias_mes = self._mini_card(sumf, "Dias", "0", COR_PRIMARY)
        for i, c in enumerate([
            self.card_trab_mes, self.card_meta_mes,
            self.card_saldo_mes, self.card_dias_mes,
        ]):
            c["frame"].grid(row=0, column=i, padx=6, pady=4, sticky="ew")
            sumf.columnconfigure(i, weight=1)

        tf = tk.Frame(f, bg=COR_CARD)
        tf.pack(fill="both", expand=True, padx=16, pady=10)
        cols = ("data", "dia", "ini", "fim", "trab", "saldo", "tipo")
        widths = (80, 65, 70, 70, 80, 90, 100)
        labels = ("Data", "Dia", "Inicio", "Fim", "Trabalho", "Saldo", "Tipo")
        tree = ttk.Treeview(tf, columns=cols, show="headings", height=11)
        for c, l, w in zip(cols, labels, widths):
            tree.heading(c, text=l)
            tree.column(c, width=w, anchor="center")
        sb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        tree.tag_configure("positivo", foreground=COR_SUCCESS)
        tree.tag_configure("negativo", foreground=COR_DANGER)
        tree.tag_configure("falta", foreground="#7a1a1a", background="#ffe3e3")
        tree.tag_configure("feriado", background="#fff4e0")
        tree.tag_configure("atestado", background="#ffe0e0")
        tree.tag_configure("folga", background="#e0f0ff")
        tree.tag_configure("extra", foreground=COR_SUCCESS, background="#e6fcf5")
        tree.tag_configure("hoje", background="#fff9db")
        tree.bind("<Double-1>", self._on_mes_doubleclick)
        self.tree_mes = tree

    def _on_mes_doubleclick(self, ev):
        sel = self.tree_mes.selection()
        if not sel:
            return
        vals = self.tree_mes.item(sel[0])["values"]
        if not vals:
            return
        hoje = date.today()
        try:
            d = datetime.strptime(f"{vals[0]}/{hoje.year}", "%d/%m/%Y").date()
        except ValueError:
            return
        self.data_vis = d
        self._carregar_dia_vis()
        self.notebook.select(0)

    # ----- aba Banco geral -----

    def _build_tab_geral(self):
        f = self.tab_geral
        hdr = tk.Frame(f, bg=COR_CARD)
        hdr.pack(fill="x", padx=16, pady=(16, 6))
        ttk.Label(hdr, text="Banco de horas geral", style="Title.TLabel").pack(side="left")

        sumf = tk.Frame(f, bg=COR_CARD)
        sumf.pack(fill="x", padx=16, pady=6)
        self.card_saldo_geral = self._mini_card(sumf, "Saldo acumulado", "00:00:00", COR_SUCCESS)
        self.card_trab_geral = self._mini_card(sumf, "Trabalhado total", "00:00:00", COR_PRIMARY)
        self.card_dias_geral = self._mini_card(sumf, "Dias trabalhados", "0", COR_PRIMARY)
        for i, c in enumerate([
            self.card_saldo_geral, self.card_trab_geral, self.card_dias_geral,
        ]):
            c["frame"].grid(row=0, column=i, padx=6, pady=4, sticky="ew")
            sumf.columnconfigure(i, weight=1)

        tf = tk.Frame(f, bg=COR_CARD)
        tf.pack(fill="both", expand=True, padx=16, pady=10)
        cols = ("mes", "dias", "trab", "meta", "saldo", "acum")
        widths = (90, 70, 110, 110, 110, 120)
        labels = ("Mes", "Dias", "Trabalhado", "Meta", "Saldo do mes", "Acumulado")
        tree = ttk.Treeview(tf, columns=cols, show="headings", height=11)
        for c, l, w in zip(cols, labels, widths):
            tree.heading(c, text=l)
            tree.column(c, width=w, anchor="center")
        sb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        tree.tag_configure("positivo", foreground=COR_SUCCESS)
        tree.tag_configure("negativo", foreground=COR_DANGER)
        self.tree_geral = tree

        ttk.Label(
            f,
            text="Saldo conta seg-sex. Dias uteis sem registro contam como -08:00. "
                 "Use FERIADO/ATESTADO/FOLGA para neutralizar.",
            style="Muted.TLabel",
        ).pack(side="bottom", pady=8)

    # ----- aba Marcacoes -----

    def _build_tab_marc(self):
        f = self.tab_marc
        hdr = tk.Frame(f, bg=COR_CARD)
        hdr.pack(fill="x", padx=16, pady=(16, 6))
        ttk.Label(
            hdr, text="Marcacoes (feriado, atestado, folga)", style="Title.TLabel"
        ).pack(side="left")

        form = tk.Frame(f, bg=COR_CARD)
        form.pack(fill="x", padx=16, pady=10)

        tk.Label(form, text="Data (DD/MM/AAAA):", bg=COR_CARD, fg=COR_TEXT,
                 font=("Segoe UI", 10)).grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self.entry_data_marc = tk.Entry(form, font=("Segoe UI", 10), width=14, relief="solid", bd=1)
        self.entry_data_marc.grid(row=0, column=1, padx=4, pady=4, sticky="w")
        self.entry_data_marc.insert(0, date.today().strftime("%d/%m/%Y"))

        tk.Label(form, text="Tipo:", bg=COR_CARD, fg=COR_TEXT,
                 font=("Segoe UI", 10)).grid(row=0, column=2, sticky="w", padx=(20, 4), pady=4)
        self.cb_tipo_marc = ttk.Combobox(
            form, values=[t for t in TIPOS_DIA if t != "NORMAL"],
            state="readonly", width=12,
        )
        self.cb_tipo_marc.grid(row=0, column=3, padx=4, pady=4, sticky="w")
        self.cb_tipo_marc.set("FERIADO")

        tk.Label(form, text="Observacao:", bg=COR_CARD, fg=COR_TEXT,
                 font=("Segoe UI", 10)).grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self.entry_obs_marc = tk.Entry(form, font=("Segoe UI", 10), width=55, relief="solid", bd=1)
        self.entry_obs_marc.grid(row=1, column=1, columnspan=3, padx=4, pady=4, sticky="ew")
        form.columnconfigure(3, weight=1)

        btnf = tk.Frame(form, bg=COR_CARD)
        btnf.grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 4))
        self._botao(btnf, "+  Salvar marcacao", COR_PRIMARY, COR_PRIMARY_HOVER, self.salvar_marcacao).grid(row=0, column=0, padx=4)
        self._botao(btnf, "✕  Remover", COR_DANGER, "#a91e1e", self.remover_marcacao_selecionada).grid(row=0, column=1, padx=4)

        tf = tk.Frame(f, bg=COR_CARD)
        tf.pack(fill="both", expand=True, padx=16, pady=10)
        cols = ("data", "tipo", "obs")
        tree = ttk.Treeview(tf, columns=cols, show="headings", height=10)
        tree.heading("data", text="Data")
        tree.heading("tipo", text="Tipo")
        tree.heading("obs", text="Observacao")
        tree.column("data", width=110, anchor="center")
        tree.column("tipo", width=110, anchor="center")
        tree.column("obs", width=420, anchor="w")
        sb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        tree.tag_configure("feriado", background="#fff4e0")
        tree.tag_configure("atestado", background="#ffe0e0")
        tree.tag_configure("folga", background="#e0f0ff")
        tree.bind("<<TreeviewSelect>>", self._on_marc_selected)
        self.tree_marc = tree

    def _on_marc_selected(self, ev):
        sel = self.tree_marc.selection()
        if not sel:
            return
        vals = self.tree_marc.item(sel[0])["values"]
        if not vals:
            return
        self.entry_data_marc.delete(0, tk.END)
        self.entry_data_marc.insert(0, vals[0])
        self.cb_tipo_marc.set(vals[1])
        self.entry_obs_marc.delete(0, tk.END)
        if len(vals) > 2 and vals[2]:
            self.entry_obs_marc.insert(0, vals[2])

    # ----- helpers UI -----

    def _botao(self, parent, texto, bg, active_bg, cmd):
        return HoverButton(
            parent, text=texto, bg=bg, hover_bg=active_bg,
            padx=18, pady=10, command=cmd,
        )

    def _mini_card(self, parent, label, valor, cor_valor):
        try:
            parent_bg = parent.cget("bg")
        except tk.TclError:
            parent_bg = COR_CARD
        cf = RoundedCard(
            parent, bg=COR_CARD, parent_bg=parent_bg,
            radius=14, shadow=True,
        )
        # altura proporcional ao conteudo - garante render correto antes do .grid
        cf.configure(height=98)
        inner = tk.Frame(cf.content, bg=COR_CARD, padx=18, pady=14)
        inner.pack(fill="both", expand=True)
        lbl_lbl = tk.Label(
            inner, text=label.upper(), bg=COR_CARD, fg=COR_MUTED,
            font=("Segoe UI", 8, "bold"), anchor="w",
        )
        lbl_lbl.pack(anchor="w")
        lbl_val = tk.Label(
            inner, text=valor, bg=COR_CARD, fg=cor_valor,
            font=("Consolas", 18, "bold"), anchor="w",
        )
        lbl_val.pack(anchor="w", pady=(4, 0))
        return {"frame": cf, "value": lbl_val, "label": lbl_lbl}

    def _set_card(self, card, valor, cor=None):
        card["value"].config(text=valor)
        if cor:
            card["value"].config(fg=cor)

    # ----- estado / persistencia do dia -----

    def _carregar_dia_vis(self):
        """Carrega self.dia_vis a partir do estado (hoje_dia se hoje) ou do Excel."""
        if self.data_vis == date.today():
            self.dia_vis = self._copia(self.hoje_dia)
        else:
            d = self.store.get_dia(self.data_vis) if self.store else None
            if d:
                self.dia_vis = {
                    "inicio": d["inicio"],
                    "ini_pausa": d["ini_pausa"],
                    "fim_pausa": d["fim_pausa"],
                    "fim": d["fim"],
                    "tipo": d["tipo"],
                    "obs": d["obs"],
                }
            else:
                self.dia_vis = self._dia_vazio()
        self._render_form()
        self._render_header_dia()

    def _copia(self, d):
        return {
            "inicio": d.get("inicio"),
            "ini_pausa": d.get("ini_pausa"),
            "fim_pausa": d.get("fim_pausa"),
            "fim": d.get("fim"),
            "tipo": d.get("tipo", "NORMAL"),
            "obs": d.get("obs", ""),
        }

    def _render_form(self):
        self.var_inicio.set(time_to_str(self.dia_vis["inicio"]) if self.dia_vis["inicio"] else "")
        self.var_ini_pausa.set(time_to_str(self.dia_vis["ini_pausa"]) if self.dia_vis["ini_pausa"] else "")
        self.var_fim_pausa.set(time_to_str(self.dia_vis["fim_pausa"]) if self.dia_vis["fim_pausa"] else "")
        self.var_fim.set(time_to_str(self.dia_vis["fim"]) if self.dia_vis["fim"] else "")
        self.var_tipo.set(self.dia_vis["tipo"] or "NORMAL")
        self.var_obs.set(self.dia_vis["obs"] or "")

    def _render_header_dia(self):
        hoje = date.today()
        eh_hoje = self.data_vis == hoje
        dia_semana = DIAS_SEMANA[self.data_vis.weekday()]
        # Linha principal: data ou "Hoje"
        if eh_hoje:
            self.lbl_data_principal.config(
                text=f"Hoje  ·  {self.data_vis.strftime('%d/%m/%Y')}",
                fg=COR_PRIMARY,
            )
        else:
            self.lbl_data_principal.config(
                text=self.data_vis.strftime("%d/%m/%Y"),
                fg=COR_TEXT,
            )
        # Sub: dia da semana por extenso + mes
        nome_mes = MESES_PT[self.data_vis.month - 1].lower()
        self.lbl_data_sub.config(
            text=f"{dia_semana} · {self.data_vis.day:02d} de {nome_mes} de {self.data_vis.year}",
        )
        self.lbl_timer_sub.config(text="Trabalhado hoje" if eh_hoje else "Trabalhado no dia")

        # Botoes live so quando eh hoje
        estado_play_pause = "normal" if eh_hoje else "disabled"
        for b in (self.btn_play, self.btn_pause, self.btn_fim_dia):
            b.config(state=estado_play_pause)
            if estado_play_pause == "disabled":
                b.config(bg=COR_DISABLED)
            else:
                # restaura cor original via HoverButton._bg_normal
                if hasattr(b, "_bg_normal"):
                    b.config(bg=b._bg_normal)

        # Botao "Hoje" vira chamativo quando NAO esta em hoje
        if eh_hoje:
            self.btn_hoje.set_palette(COR_PRIMARY_LIGHT, "#c7d2fe")
            self.btn_hoje.config(fg=COR_PRIMARY_TEXT, text="● Hoje")
        else:
            self.btn_hoje.set_palette(COR_PRIMARY, COR_PRIMARY_HOVER)
            self.btn_hoje.config(fg="white", text="↺ Voltar pra hoje")

        existe = self.store and self.store.get_dia(self.data_vis) is not None
        self.btn_excluir.config(state="normal" if existe else "disabled")
        if existe:
            self.btn_excluir.set_palette(COR_CARD, COR_DANGER_LIGHT)
            self.btn_excluir.config(fg=COR_DANGER)
        else:
            self.btn_excluir.config(fg=COR_DISABLED, bg=COR_CARD)

    def _commit_campo(self, key):
        var_map = {
            "inicio": self.var_inicio,
            "ini_pausa": self.var_ini_pausa,
            "fim_pausa": self.var_fim_pausa,
            "fim": self.var_fim,
        }
        var = var_map[key]
        valor_str = var.get().strip()
        if not valor_str:
            self.dia_vis[key] = None
        else:
            t = parse_hhmmss(valor_str)
            if t is None:
                messagebox.showerror(
                    "Hora invalida",
                    f"Use HH:MM ou HH:MM:SS. Valor recebido: {valor_str!r}"
                )
                # restaura
                self._render_form()
                return
            self.dia_vis[key] = datetime.combine(self.data_vis, t)
        # Reflete no estado de hoje
        if self.data_vis == date.today():
            self.hoje_dia[key] = self.dia_vis[key]
        # Re-renderiza pra formatar HH:MM:SS
        self._render_form()
        # Atualiza cards
        self._atualizar_timer_para_dia_vis()

    def _on_tipo_change(self):
        self.dia_vis["tipo"] = self.var_tipo.get()
        if self.data_vis == date.today():
            self.hoje_dia["tipo"] = self.dia_vis["tipo"]

    def _capturar_obs(self):
        self.dia_vis["obs"] = self.var_obs.get()
        if self.data_vis == date.today():
            self.hoje_dia["obs"] = self.dia_vis["obs"]

    # ----- navegacao -----

    def navegar_dia(self, delta):
        nova = self.data_vis + timedelta(days=delta)
        if nova > date.today():
            return
        self.data_vis = nova
        self._carregar_dia_vis()
        self._atualizar_timer_para_dia_vis()

    def ir_para_hoje(self):
        self.data_vis = date.today()
        self._carregar_dia_vis()
        self._atualizar_timer_para_dia_vis()

    # ----- play / pause / finalizar -----

    def _registrar_evento(self, descricao):
        hora = datetime.now().strftime("%H:%M:%S")
        self.hoje_eventos.append(f"[{hora}] {descricao}")

    def play(self):
        if self.data_vis != date.today():
            return
        agora = datetime.now()
        d = self.hoje_dia
        if d["ini_pausa"] and not d["fim_pausa"]:
            # estava em pausa -> retoma
            d["fim_pausa"] = agora
            self._registrar_evento("FIM DA PAUSA")
        elif not d["inicio"]:
            d["inicio"] = agora
            self._registrar_evento("INÍCIO DO TRABALHO")
        # se ja trabalhando sem pausa, play eh no-op
        self.dia_vis = self._copia(d)
        self._render_form()

    def pause(self):
        if self.data_vis != date.today():
            return
        agora = datetime.now()
        d = self.hoje_dia
        if d["inicio"] and not d["ini_pausa"]:
            d["ini_pausa"] = agora
            self._registrar_evento("INÍCIO DA PAUSA")
        elif d["fim_pausa"] and not d["fim"]:
            # ja teve pausa; segunda pausa nao suportada pelo modelo
            messagebox.showinfo(
                "Limite de pausa",
                "Este app modela 1 pausa por dia. Edite os horarios manualmente "
                "se precisar registrar varias pausas."
            )
        self.dia_vis = self._copia(d)
        self._render_form()

    def finalizar_dia(self):
        if self.data_vis != date.today():
            return
        d = self.hoje_dia
        if not d["inicio"]:
            messagebox.showinfo("Sem inicio", "Aperte Play primeiro ou preencha o Inicio.")
            return
        if not messagebox.askyesno("Finalizar dia", "Confirma finalizacao do dia?"):
            return
        agora = datetime.now()
        if d["ini_pausa"] and not d["fim_pausa"]:
            d["fim_pausa"] = agora
            self._registrar_evento("FIM DA PAUSA")
            self._registrar_evento("FINALIZAÇÃO DO DIA (EM PAUSA)")
        else:
            self._registrar_evento("FINALIZAÇÃO DO DIA")
        d["fim"] = agora
        self.dia_vis = self._copia(d)
        self._render_form()
        self._gravar_txt()
        ok_excel = self._persistir_dia_no_excel(date.today())
        trab, _ = calcular_trab_pausa(d["inicio"], d["ini_pausa"], d["fim_pausa"], d["fim"])
        saldo = trab - META_DIARIA
        msg = (
            f"Dia salvo!\n\n"
            f"Trabalho: {td_to_str(trab)}\n"
            f"Meta:     08:00:00\n"
            f"Saldo:    {td_to_str(saldo)}\n\n"
            f"Excel:    {'OK' if ok_excel else 'NAO gravado'}\n"
            f"TXT:      ponto_{agora.strftime('%Y_%m')}.txt"
        )
        # reset estado pra novo dia (mas mantem visivel)
        self.hoje_eventos = []
        self.atualizar_visoes()
        messagebox.showinfo("Pronto", msg)

    def _gravar_txt(self):
        if not self.hoje_eventos:
            return
        agora = datetime.now()
        nome_txt = DATA_DIR / agora.strftime("ponto_%Y_%m.txt")
        d = self.hoje_dia
        trab, pausa = calcular_trab_pausa(d["inicio"], d["ini_pausa"], d["fim_pausa"], d["fim"])
        try:
            with open(nome_txt, "a", encoding="utf-8") as f:
                f.write(f"{'=' * 30}\n")
                f.write(f"Data: {agora.strftime('%d/%m/%Y')}\n\n")
                for evento in self.hoje_eventos:
                    f.write(evento + "\n")
                f.write("\nResumo:\n")
                f.write(f"Trabalho total: {td_to_str(trab)}\n")
                f.write(f"Pausa total:    {td_to_str(pausa)}\n")
                f.write(f"{'=' * 30}\n\n")
        except Exception as e:
            messagebox.showwarning("Erro no .txt", f"Nao gravou {nome_txt.name}:\n{e}")

    # ----- salvar / excluir manualmente -----

    def salvar_dia_atual(self):
        # Garante que os ultimos textos digitados foram aplicados
        for key in ("inicio", "ini_pausa", "fim_pausa", "fim"):
            self._commit_campo_se_dirty(key)
        self._capturar_obs()
        if not self.store:
            messagebox.showerror("Sem Excel", "openpyxl nao instalado.")
            return
        # Se hoje, sincroniza
        if self.data_vis == date.today():
            self.hoje_dia = self._copia(self.dia_vis)
        ok = self._persistir_dia_no_excel(self.data_vis)
        if ok:
            mostrar_toast(
                self.root,
                f"✓  {self.data_vis.strftime('%d/%m/%Y')} salvo",
                cor=COR_SUCCESS,
            )
        self.atualizar_visoes()
        self._render_header_dia()

    def _commit_campo_se_dirty(self, key):
        var_map = {
            "inicio": self.var_inicio,
            "ini_pausa": self.var_ini_pausa,
            "fim_pausa": self.var_fim_pausa,
            "fim": self.var_fim,
        }
        novo = var_map[key].get().strip()
        atual = time_to_str(self.dia_vis[key]) if self.dia_vis[key] else ""
        if novo != atual:
            self._commit_campo(key)

    def _persistir_dia_no_excel(self, data):
        if not self.store:
            return False
        d = self.dia_vis if data == self.data_vis else None
        if d is None:
            return False
        try:
            self.store.salvar_dia(
                data, d["inicio"], d["ini_pausa"], d["fim_pausa"], d["fim"],
                d.get("tipo", "NORMAL"), d.get("obs", ""),
            )
            # sincroniza marcacao
            tipo = d.get("tipo", "NORMAL")
            if tipo != "NORMAL":
                self.store.set_marcacao(data, tipo, d.get("obs", ""))
            else:
                # se tinha marcacao, remove
                if data in self.store.listar_marcacoes():
                    self.store.remover_marcacao(data)
            return True
        except PermissionError:
            messagebox.showerror(
                "Excel aberto",
                f"Feche {EXCEL_PATH.name} no Excel e tente novamente."
            )
            return False
        except Exception as e:
            messagebox.showerror("Erro ao salvar", str(e))
            return False

    def excluir_dia_atual(self):
        if not self.store:
            return
        d = self.store.get_dia(self.data_vis)
        if not d:
            messagebox.showinfo("Nada para excluir", "Esse dia nao tem registro no Excel.")
            return
        if not messagebox.askyesno(
            "Excluir dia",
            f"Excluir registro de {self.data_vis.strftime('%d/%m/%Y')} do Excel?\n"
            "Isto NAO apaga o .txt."
        ):
            return
        try:
            self.store.excluir_dia(self.data_vis)
            if self.data_vis == date.today():
                self.hoje_dia = self._dia_vazio()
                self.hoje_eventos = []
            self.dia_vis = self._dia_vazio()
            self._render_form()
            self.atualizar_visoes()
            self._render_header_dia()
        except PermissionError:
            messagebox.showerror(
                "Excel aberto",
                f"Feche {EXCEL_PATH.name} no Excel e tente novamente."
            )

    # ----- marcacoes UI -----

    def salvar_marcacao(self):
        if not self.store:
            return
        try:
            d = datetime.strptime(self.entry_data_marc.get().strip(), "%d/%m/%Y").date()
        except ValueError:
            messagebox.showerror("Data invalida", "Use o formato DD/MM/AAAA.")
            return
        tipo = self.cb_tipo_marc.get()
        try:
            self.store.set_marcacao(d, tipo, self.entry_obs_marc.get().strip())
        except PermissionError:
            messagebox.showerror("Excel aberto",
                                 f"Feche {EXCEL_PATH.name} no Excel e tente novamente.")
            return
        mostrar_toast(self.root, f"✓  {tipo} em {d.strftime('%d/%m')}", cor=COR_SUCCESS)
        self.atualizar_visoes()
        if d == self.data_vis:
            self._carregar_dia_vis()

    def remover_marcacao_selecionada(self):
        if not self.store:
            return
        sel = self.tree_marc.selection()
        if not sel:
            messagebox.showinfo("Nada selecionado", "Selecione uma marcacao para remover.")
            return
        vals = self.tree_marc.item(sel[0])["values"]
        if not vals:
            return
        try:
            d = datetime.strptime(vals[0], "%d/%m/%Y").date()
        except ValueError:
            return
        if not messagebox.askyesno("Confirmar", f"Remover marcacao de {vals[0]}?"):
            return
        try:
            self.store.remover_marcacao(d)
        except PermissionError:
            messagebox.showerror("Excel aberto",
                                 f"Feche {EXCEL_PATH.name} no Excel e tente novamente.")
            return
        self.atualizar_visoes()

    # ----- cronometro live -----

    def atualizar_tempo(self):
        self._atualizar_timer_para_dia_vis()
        self.root.after(1000, self.atualizar_tempo)

    def _autosave_hoje(self):
        try:
            d = self.hoje_dia
            if self.store and d.get("inicio") and not d.get("fim"):
                self.store.salvar_dia(
                    date.today(),
                    d["inicio"], d["ini_pausa"], d["fim_pausa"], d["fim"],
                    d.get("tipo", "NORMAL"), d.get("obs", ""),
                )
        except Exception as e:
            print(f"Auto-save falhou: {e}")
        finally:
            self.root.after(AUTOSAVE_INTERVALO_MS, self._autosave_hoje)

    def _atualizar_timer_para_dia_vis(self):
        d = self.dia_vis
        agora = datetime.now() if self.data_vis == date.today() else None
        trab, pausa = calcular_trab_pausa(
            d["inicio"], d["ini_pausa"], d["fim_pausa"], d["fim"], agora=agora,
        )
        # Cor do timer muda conforme % da meta
        meta_s = META_DIARIA.total_seconds() or 1
        pct = max(0.0, min(100.0, (trab.total_seconds() / meta_s) * 100))
        if pct >= 100:
            cor_timer = COR_SUCCESS
        elif pct >= 75:
            cor_timer = COR_SUCCESS
        elif pct >= 40:
            cor_timer = COR_PRIMARY
        elif pct > 0:
            cor_timer = COR_WARNING
        else:
            cor_timer = COR_MUTED
        self.lbl_timer_big.config(text=td_to_str(trab), fg=cor_timer)

        self._set_card(self.card_pausa_dia, td_to_str(pausa))
        saldo = trab - META_DIARIA if d["inicio"] else timedelta()
        self._set_card(
            self.card_saldo_dia,
            td_to_str(saldo),
            COR_SUCCESS if saldo.total_seconds() >= 0 else COR_DANGER,
        )
        self.progress["value"] = pct
        self.lbl_pct.config(text=f"{pct:.1f}% da meta diaria")

        # status pill
        if self.data_vis == date.today():
            if d["inicio"] and not d["fim"] and d["ini_pausa"] and not d["fim_pausa"]:
                self._set_pill(self.card_status_dia, "Em pausa", COR_WARNING)
            elif d["inicio"] and not d["fim"]:
                self._set_pill(self.card_status_dia, "Trabalhando", COR_SUCCESS)
            elif d["fim"]:
                self._set_pill(self.card_status_dia, "Finalizado", COR_PRIMARY)
            else:
                self._set_pill(self.card_status_dia, "Parado", COR_MUTED)
        else:
            if d["inicio"] and d["fim"]:
                self._set_pill(self.card_status_dia, "Visualizando", COR_PRIMARY)
            elif d["inicio"]:
                self._set_pill(self.card_status_dia, "Parcial", COR_WARNING)
            else:
                self._set_pill(self.card_status_dia, "Sem registro", COR_MUTED)

    def _set_pill(self, card, texto, cor):
        card["value"].config(text=texto, bg=cor)
        card["pill"].config(bg=cor)

    # ----- views agregadas -----

    def atualizar_visoes(self):
        if not self.store:
            return
        try:
            saldo_mes, saldo_geral, entradas_mes = calcular_saldos(self.store)
        except Exception as e:
            print(f"Erro calcular_saldos: {e}")
            return
        hoje = date.today()

        if hasattr(self, "tree_mes"):
            self.tree_mes.delete(*self.tree_mes.get_children())
            soma_trab = timedelta()
            soma_dias = 0
            mapa_dias_excel = {}
            for d in self.store.listar_dias_do_mes(hoje.year, hoje.month):
                mapa_dias_excel[d["data"]] = d
            for e in entradas_mes:
                tipo = e["tipo"]
                tags = []
                if tipo == "NORMAL":
                    if e["saldo"].total_seconds() > 0:
                        tags.append("positivo")
                    elif e["saldo"].total_seconds() < 0:
                        tags.append("negativo")
                    soma_trab += e["trabalho"]
                    soma_dias += 1
                elif tipo == "FALTA":
                    tags.append("falta")
                elif tipo == "FERIADO":
                    tags.append("feriado")
                    soma_trab += META_DIARIA
                elif tipo == "ATESTADO":
                    tags.append("atestado")
                    soma_trab += META_DIARIA
                elif tipo == "FOLGA":
                    tags.append("folga")
                    soma_trab += META_DIARIA
                elif tipo == "EXTRA":
                    tags.append("extra")
                if e["data"] == hoje:
                    tags.append("hoje")
                trab_str = td_to_str(e["trabalho"]) if e["trabalho"].total_seconds() > 0 else "-"
                ini = "-"
                fim = "-"
                dd = mapa_dias_excel.get(e["data"])
                if dd:
                    ini = time_to_str(dd["inicio"]) if dd["inicio"] else "-"
                    fim = time_to_str(dd["fim"]) if dd["fim"] else "-"
                self.tree_mes.insert("", "end", values=(
                    e["data"].strftime("%d/%m"),
                    e["dia_semana"][:3],
                    ini, fim,
                    trab_str,
                    td_to_str(e["saldo"]),
                    tipo,
                ), tags=tags)
            soma_meta = META_DIARIA * dias_uteis_no_mes(hoje.year, hoje.month)
            self._set_card(self.card_trab_mes, td_to_str(soma_trab))
            self._set_card(self.card_meta_mes, td_to_str(soma_meta))
            self._set_card(
                self.card_saldo_mes,
                td_to_str(saldo_mes),
                COR_SUCCESS if saldo_mes.total_seconds() >= 0 else COR_DANGER,
            )
            self._set_card(self.card_dias_mes, str(soma_dias))

        if hasattr(self, "tree_geral"):
            self.tree_geral.delete(*self.tree_geral.get_children())
            meses = self.store.listar_meses()
            acumulado = timedelta()
            total_trab = timedelta()
            total_dias = 0
            for a, m in meses:
                dias = self.store.listar_dias_do_mes(a, m)
                normais = [d for d in dias if d["tipo"] == "NORMAL"]
                neutros = [d for d in dias if d["tipo"] in TIPOS_NEUTRALIZA_META]
                trab = sum((d["trabalho"] for d in normais), timedelta())
                trab += META_DIARIA * len(neutros)
                meta = META_DIARIA * dias_uteis_no_mes(a, m)
                saldo_m = trab - meta
                acumulado += saldo_m
                total_trab += trab
                total_dias += len(normais)
                tags = []
                if saldo_m.total_seconds() > 0:
                    tags.append("positivo")
                elif saldo_m.total_seconds() < 0:
                    tags.append("negativo")
                self.tree_geral.insert("", "end", values=(
                    f"{a:04d}-{m:02d}",
                    str(len(normais)),
                    td_to_str(trab), td_to_str(meta),
                    td_to_str(saldo_m), td_to_str(acumulado),
                ), tags=tags)
            self._set_card(
                self.card_saldo_geral,
                td_to_str(saldo_geral),
                COR_SUCCESS if saldo_geral.total_seconds() >= 0 else COR_DANGER,
            )
            self._set_card(self.card_trab_geral, td_to_str(total_trab))
            self._set_card(self.card_dias_geral, str(total_dias))

        if hasattr(self, "tree_marc"):
            self.tree_marc.delete(*self.tree_marc.get_children())
            marc = self.store.listar_marcacoes()
            for dt in sorted(marc.keys(), reverse=True):
                tipo, obs = marc[dt]
                tag = tipo.lower() if tipo.lower() in ("feriado", "atestado", "folga") else ""
                self.tree_marc.insert("", "end", values=(
                    dt.strftime("%d/%m/%Y"), tipo, obs,
                ), tags=(tag,) if tag else ())


if __name__ == "__main__":
    # AppUserModelID precisa ser setado ANTES de criar a janela para que a
    # barra de tarefas do Windows associe o icone do app (em vez do icone
    # do interpretador Python).
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
            APP_USER_MODEL_ID
        )
    except Exception:
        pass
    root = tk.Tk()
    app = FolhaPontoApp(root)
    root.mainloop()
